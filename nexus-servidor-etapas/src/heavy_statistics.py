"""Independent SENATRAN type/species/axles statistics.

These rows are deliberately never joined to model observations. They describe
the composition of the heavy fleet, not the identity or propulsion of a model.
"""
import hashlib,re,uuid,zipfile
from pathlib import Path
import pandas as pd
from sqlalchemy import select,insert,update,func
from src import db,intelligence_schema as s,operational_schema as o,pipeline_schema as ps
from src.matching import normalize
from src.transform.ingest import csv_frames,UF_NAMES

HEAVY_TYPES={
    'ONIBUS':'bus','MICRO ONIBUS':'microbus','CAMINHAO':'truck','CAMINHAO TRATOR':'tractor_truck',
    'UTILITARIO':'utility','CAMINHONETE':'utility','CAMIONETA':'utility','CHASSI PLATAFORMA':'chassis_platform',
    'REBOQUE':'trailer','SEMI REBOQUE':'semitrailer'
}
ALIASES={'UF':'state','MUNICIPIO':'municipality','MUNICIPIO NOME':'municipality','TIPO VEICULO':'vehicle_type',
    'TIPO DE VEICULO':'vehicle_type','ESPECIE':'species','EIXOS':'axles','QUANTIDADE EIXOS':'axles',
    'QUANTIDADE':'quantity','QTD VEICULOS':'quantity','QTD':'quantity','QTDE':'quantity','TOTAL':'quantity'}

def stable(value):return str(uuid.uuid5(uuid.NAMESPACE_URL,'i9-heavy-stat:'+value))

def frames(path):
    is_xlsx=False
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as probe:is_xlsx='xl/workbook.xml' in probe.namelist()
    if is_xlsx:
        from openpyxl import load_workbook
        source=open(path,'rb');workbook=load_workbook(source,read_only=True,data_only=True)
        try:
            for sheet in workbook.worksheets:
                iterator=sheet.iter_rows(values_only=True)
                header=next(iterator,None)
                if not header:continue
                columns=[str(value or '').strip() for value in header]
                chunk=[]
                for values in iterator:
                    chunk.append(values)
                    if len(chunk)>=100_000:
                        yield sheet.title,pd.DataFrame(chunk,columns=columns);chunk=[]
                if chunk:yield sheet.title,pd.DataFrame(chunk,columns=columns)
        finally:
            workbook.close();source.close()
    elif zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as archive:
            members=[m for m in archive.infolist() if not m.is_dir() and m.filename.lower().endswith(('.csv','.txt'))]
            if not members:raise ValueError('Arquivo 08 sem CSV/TXT')
            if sum(m.file_size for m in members)>8_000_000_000:raise ValueError('Arquivo 08 excede 8 GB descompactados')
            for member in sorted(members,key=lambda m:m.filename):
                with archive.open(member) as stream:
                    for frame in csv_frames(stream):yield member.filename,frame
    else:
        with open(path,'rb') as stream:
            for frame in csv_frames(stream):yield Path(path).name,frame

def import_type_species_axles(path,reference,lease,url=None):
    from src.pipeline import _block_result,_commit_block,_upsert_nothing,job_transaction,lock_lease,timestamp,canonical
    path=Path(path).resolve()
    with path.open('rb') as stream:raw_hash=hashlib.file_digest(stream,'sha256').hexdigest()
    source_id=stable(f'SENATRAN|tipo-especie-eixos|{reference}|{raw_hash}')
    with job_transaction(lease,'heavy-artifact') as c:
        _upsert_nothing(c,s.source_snapshot,dict(id=source_id,provider='SENATRAN',dataset_name='tipo-especie-eixos',
            reference_period=reference,retrieved_at=timestamp(),method='public_monthly_aggregate',raw_file_hash=raw_hash,
            raw_storage_path=None,source_version='2026',transformation_version='heavy-stat-1',license_or_terms='Dados públicos SENATRAN',
            ingestion_status='staging',source_url=url,scope='Estatística independente de tipo, espécie e eixos',metric='fleet_stock',
            access_level='public',record_count=0,metadata={'pipeline_job_id':lease['job_id'],'no_model_join':True,
                'scope':'metadata_only; tipo/espécie/eixos não informa eletrificação'}))
    # This report contains combustion vehicles and cannot be joined to the
    # model-level electrification evidence. Keep provenance only.
    with job_transaction(lease,'heavy-complete') as c:
        c.execute(update(s.source_snapshot).where(s.source_snapshot.c.id==source_id).values(ingestion_status='success',record_count=0))
    if path.is_relative_to((db.ROOT/'data/raw/pipeline').resolve()):path.unlink(missing_ok=True)
    return {'source_snapshot_id':source_id,'reference':reference,'rows':0,'dataset':'tipo-especie-eixos','no_model_join':True}
