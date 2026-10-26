"""Bounded import of authorized CSV/XLSX exports, with an explicit column contract."""
import csv
import io
import zipfile
from src.intelligence_models import Observation

def parse_tabular(raw,format):
    if len(raw)>10_000_000:raise ValueError('Arquivo excede 10 MB')
    if format=='csv':
        text=raw.decode('utf-8-sig')
        rows=list(csv.reader(io.StringIO(text),delimiter=';' if ';' in text.splitlines()[0] else ','))
    elif format=='xlsx':
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if sum(x.file_size for x in z.infolist())>50_000_000:raise ValueError('Planilha expandida excede o limite')
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(raw),read_only=True,data_only=False)
        try:
            rows=[]
            for i,row in enumerate(book.active.iter_rows(values_only=True)):
                if i>10000:raise ValueError('Limite de 10.000 observações')
                rows.append(list(row))
        finally:book.close()
    else:raise ValueError('Formato deve ser csv ou xlsx')
    if not rows:raise ValueError('Arquivo vazio')
    headers=[str(x).strip() for x in rows[0]]
    if len(headers)!=len(set(headers)) or set(headers)-set(Observation.model_fields):raise ValueError('Cabeçalhos duplicados ou não permitidos; use o modelo de importação')
    if not {'quantity','source_record_reference'}<=set(headers):raise ValueError('quantity e source_record_reference são obrigatórios')
    out=[]
    for row in rows[1:]:
        if all(x is None or x=='' for x in row):continue
        if len(row)!=len(headers):raise ValueError('Número de colunas inconsistente')
        record={k:v for k,v in zip(headers,row) if v is not None and v!=''}
        for k in ('quantity','model_year','manufacture_year'):
            if k in record:
                v=str(record[k]).strip()
                if not v.isdigit():raise ValueError(f'{k} deve conter inteiro não negativo')
                record[k]=int(v)
        if 'source_record_reference' in record:record['source_record_reference']=str(record['source_record_reference'])
        out.append(Observation.model_validate(record))
        if len(out)>10000:raise ValueError('Limite de 10.000 observações')
    if not out:raise ValueError('Sem observações no arquivo')
    return out
